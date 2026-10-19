"""
Downloading a sample resume as Word or PDF: the files match what the page
shows, and no page ends on a heading cut off from the lines under it.
"""
import io
import re
import unittest
import zipfile

import helpers  # noqa: F401

from resume_builder import resume_file as rf

try:
    import docx  # noqa: F401
    import fpdf  # noqa: F401
    HAVE_LIBS = True
except ImportError:
    HAVE_LIBS = False


def resume(jobs=7, lines=5):
    long = ("that runs long enough to wrap onto a second line in the layout, "
            "with a figure like 35% and a dash — and “quotes”.")
    return {"name": "Ann Lee", "headline": "Analytics leader", "contact": "Hummelstown, PA",
            "summary": "Twenty years of analytics.",
            "employers": [{"employer": f"Employer {i}", "roles": [
                {"title": f"Director {i}", "dates": "2019 - Present",
                 "lines": [f"Line {j} for job {i} {long}" for j in range(lines)]}]}
                for i in range(jobs)],
            "education": ["MBA, Penn State"], "skills": ["SQL", "Python"],
            "extras": [{"heading": "Awards", "lines": ["Best paper"]}]}


class Paragraphs(unittest.TestCase):

    def test_headings_titles_and_first_line_stay_with_what_follows(self):
        p = rf.paragraphs(resume(jobs=1, lines=3))
        keep = {(k, c if isinstance(c, str) else c[0]): kw for k, c, kw in p}
        self.assertTrue(keep[("heading", "EXPERIENCE")])
        self.assertTrue(keep[("employer", "Employer 0")])
        self.assertTrue(keep[("role", "Director 0")])
        bullets = [kw for k, _, kw in p if k == "bullet"]
        # The first line holds on to the second; the rest may break freely.
        self.assertEqual(bullets, [True, False, False])

    def test_a_title_with_no_lines_does_not_hold_on_to_the_next_job(self):
        r = resume(jobs=2, lines=0)
        roles = [kw for k, _, kw in rf.paragraphs(r) if k == "role"]
        self.assertEqual(roles, [False, False])


@unittest.skipUnless(HAVE_LIBS, "python-docx and fpdf2 not installed")
class Files(unittest.TestCase):

    def pages(self, data):
        """The text of each PDF page, as the lines a reader would see."""
        from fpdf import FPDF  # noqa: F401  (only to be sure the lib is real)
        import subprocess
        try:
            out = subprocess.run(["pdftotext", "-layout", "-", "-"], input=data,
                                 capture_output=True, check=True).stdout.decode()
        except (OSError, subprocess.CalledProcessError):
            self.skipTest("pdftotext not available")
        return [[l.strip() for l in page.splitlines() if l.strip()] for page in out.split("\f")]

    def test_no_pdf_page_ends_on_an_employer_or_job_title(self):
        for jobs, lines in ((7, 5), (9, 3), (12, 2), (6, 7)):
            pages = [p for p in self.pages(rf.to_pdf(resume(jobs, lines))) if p]
            self.assertGreater(len(pages), 1)
            for page in pages[:-1]:
                last = page[-1]
                self.assertFalse(re.match(r"^(Employer|Director) \d+", last), (jobs, lines, last))
                self.assertNotIn(last, ("EXPERIENCE", "SKILLS", "SUMMARY"))

    def test_the_pdf_holds_unicode_text(self):
        text = " ".join(" ".join(p) for p in self.pages(rf.to_pdf(resume(1, 1))))
        self.assertIn("“quotes”", text)
        self.assertIn("ANN LEE", text)

    def test_the_word_file_keeps_headings_with_the_next_line(self):
        from docx import Document
        doc = Document(io.BytesIO(rf.to_docx(resume(2, 3))))
        by_text = {p.text: p for p in doc.paragraphs}
        self.assertTrue(by_text["EXPERIENCE"].paragraph_format.keep_with_next)
        self.assertTrue(by_text["Employer 1"].paragraph_format.keep_with_next)
        self.assertTrue(by_text["Director 1\t2019 - Present"].paragraph_format.keep_with_next)
        self.assertTrue(all(p.paragraph_format.keep_together for p in doc.paragraphs))
        self.assertEqual(doc.sections[0].page_width.inches, 8.5)

    def test_the_word_file_is_a_real_docx(self):
        data = rf.to_docx(resume(1, 1))
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.assertIn("word/document.xml", z.namelist())


if __name__ == "__main__":
    unittest.main()
