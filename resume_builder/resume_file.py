"""
A sample resume as a file to download: Word (.docx) or PDF, laid out the
way the page shows it.

Both are built from one list of paragraphs, each marked with whether it must
stay on the same page as the next. That one rule keeps a page from ending on
a stranded heading: a section heading stays with what follows it, and an
employer's name and job title stay with at least the first two lines under
them, so a break that would split them moves above the employer instead.
No single line is split across pages. A whole employer is not forced onto
one page, which would leave large gaps.

Word gets the rule as "keep with next" and "keep lines together", so it
still holds after the person edits the file.
"""
import io
from pathlib import Path

FONTS = Path(__file__).parent / "fonts"
WITH_FIRST_LINES = 2   # lines that stay with their job title


# ---------------------------------------------------------------- the paragraphs

def paragraphs(res: dict) -> list:
    """[(kind, content, keep_with_next)] in reading order."""
    out = []
    add = lambda kind, content, keep=False: out.append((kind, content, keep))
    add("name", (res.get("name") or "Your name").upper())
    line = " · ".join(x for x in (res.get("headline"), res.get("contact")) if x)
    if line:
        add("contact", line)

    def section(heading, items):
        """items: [(kind, content)]; the heading stays with the first."""
        if items:
            add("heading", heading, True)
            for kind, content in items:
                add(kind, content)

    if res.get("summary"):
        section("SUMMARY", [("text", res["summary"])])
    exp = []
    for e in res.get("employers", []):
        exp.append(("employer", e["employer"]))
        for role in e.get("roles", []):
            exp.append(("role", (role.get("title", ""), role.get("dates", ""))))
            exp.extend(("bullet", l) for l in role.get("lines", []))
    if exp:
        add("heading", "EXPERIENCE", True)
        under = 0   # lines so far under the current job title
        for n, (kind, content) in enumerate(exp):
            next_is_line = n + 1 < len(exp) and exp[n + 1][0] == "bullet"
            if kind == "employer":
                keep = True
            elif kind == "role":
                keep, under = next_is_line, 0
            else:
                under += 1
                keep = next_is_line and under < WITH_FIRST_LINES
            add(kind, content, keep)
    section("EDUCATION AND LICENCES", [("text", x) for x in res.get("education", [])])
    if res.get("skills"):
        section("SKILLS", [("text", " · ".join(res["skills"]))])
    for x in res.get("extras", []):
        section(x["heading"].upper(), [("text", l) for l in x.get("lines", [])])
    return out


# ---------------------------------------------------------------- Word

def to_docx(res: dict) -> bytes:
    return _docx(paragraphs(res))


def _docx(items: list) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.7))
    width = sec.page_width - sec.left_margin - sec.right_margin
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    muted = RGBColor(0x4A, 0x57, 0x54)

    def rule_above(p):
        border = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", "16201E")):
            top.set(qn(f"w:{k}"), v)
        border.append(top)
        p._p.get_or_add_pPr().append(border)

    for kind, content, keep in items:
        p = doc.add_paragraph()
        f = p.paragraph_format
        f.keep_with_next, f.keep_together = keep, True
        f.space_before, f.space_after = Pt(0), Pt(2)
        if kind == "name":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(content); r.bold = True; r.font.size = Pt(16)
        elif kind in ("contact", "note"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(content); r.font.color.rgb = muted
            if kind == "note":
                r.italic, r.font.size = True, Pt(9)
            f.space_after = Pt(4)
        elif kind == "heading":
            rule_above(p)
            f.space_before = Pt(8)
            r = p.add_run(content); r.bold = True; r.font.size = Pt(10)
        elif kind == "employer":
            f.space_before = Pt(5)
            p.add_run(content).bold = True
        elif kind == "role":
            title, dates = content
            f.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(title).italic = True
            if dates:
                d = p.add_run("\t" + dates); d.font.color.rgb = muted
        elif kind == "field":
            label, value = content
            p.add_run(f"{label}: ").bold = True
            p.add_run(value)
        elif kind == "subhead":
            f.space_before = Pt(4)
            r = p.add_run(content); r.italic = True; r.font.color.rgb = muted
        elif kind == "bullet":
            p.style = doc.styles["List Bullet"]
            p.add_run(content)
        elif kind == "detail":
            label, value = content
            f.left_indent = Inches(0.5)
            r = p.add_run(f"{label}: "); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = muted
            r = p.add_run(value); r.font.size = Pt(9.5); r.font.color.rgb = muted
        else:
            p.add_run(content)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def to_pdf(res: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(unit="pt", format="letter")
    margin = 0.7 * 72
    pdf.set_margins(margin, margin, margin)
    pdf.set_auto_page_break(False)
    for style, file in (("", "Regular"), ("B", "Bold"), ("I", "Italic")):
        pdf.add_font("Body", style, str(FONTS / f"LiberationSans-{file}.ttf"))
    pdf.add_page()
    width = pdf.w - 2 * margin
    bottom = pdf.h - margin
    size, lead = 10.5, 13.5
    ink, grey = (22, 32, 30), (74, 87, 84)

    def text(t, w, x, style="", fsize=size, align="L", color=ink, dry=False):
        pdf.set_font("Body", style, fsize)
        pdf.set_text_color(*color)
        pdf.set_x(x)
        return pdf.multi_cell(w, lead * fsize / size, t, align=align, dry_run=dry,
                              output="HEIGHT", new_x="LMARGIN", new_y="NEXT")

    def draw(kind, content, dry):
        """Draws one paragraph (or only measures it); returns its height."""
        x, y = margin, pdf.get_y()
        if kind == "name":
            return text(content, width, x, "B", 16, "C", dry=dry) + 2
        if kind == "contact":
            return text(content, width, x, align="C", color=grey, dry=dry) + 4
        if kind == "heading":
            if not dry:
                pdf.set_y(y + 8)
                pdf.set_draw_color(*ink)
                pdf.line(margin, y + 8, margin + width, y + 8)
                pdf.set_y(y + 12)
            return 12 + text(content, width, x, "B", 10, dry=dry)
        if kind == "employer":
            if not dry:
                pdf.set_y(y + 5)
            return 5 + text(content, width, x, "B", dry=dry)
        if kind == "role":
            title, dates = content
            pdf.set_font("Body", "", size)
            dw = pdf.get_string_width(dates) + 6 if dates else 0
            if not dry and dates:
                pdf.set_xy(margin + width - dw, y)
                pdf.set_text_color(*grey)
                pdf.cell(dw, lead, dates, align="R")
                pdf.set_xy(margin, y)
            return text(title, width - dw, x, "I", dry=dry)
        if kind == "bullet":
            indent = 14
            if not dry:
                pdf.set_font("Body", "", size)
                pdf.set_text_color(*ink)
                pdf.set_xy(margin + 3, y)
                pdf.cell(indent - 3, lead, "•")
                pdf.set_xy(margin + indent, y)
            return text(content, width - indent, margin + indent, dry=dry) + 1
        return text(content, width, x, dry=dry) + 1

    items = paragraphs(res)
    for n, (kind, content, keep) in enumerate(items):
        starts_group = n == 0 or not items[n - 1][2]
        if starts_group:
            # This paragraph and every one glued to it must fit, or the whole
            # group moves to the next page (unless it already starts one).
            need, m = 0, n
            while m < len(items):
                need += draw(items[m][0], items[m][1], True)
                if not items[m][2]:
                    break
                m += 1
            if pdf.get_y() + need > bottom and pdf.get_y() > margin + 1:
                pdf.add_page()
        elif pdf.get_y() + draw(kind, content, True) > bottom:
            pdf.add_page()
        draw(kind, content, False)
    return bytes(pdf.output())


# ---------------------------------------------------------------- the master document

def master_paragraphs(rec) -> list:
    """The whole record, every job and every line with the story behind it,
    as [(kind, content, keep_with_next)]. No page limit: this is the copy the
    person keeps, not a resume."""
    from . import journey, skills as sk
    from . import record as mr

    out = []
    add = lambda kind, content, keep=False: out.append((kind, content, keep))
    c = {k: journey.clean(rec.contact.get(k, "")) for k in journey.CONTACT_FIELDS}
    add("name", (c["Name"] or "Your career record").upper())
    line = " · ".join(x for x in (c["Location"], c["Email"], c["Telephone"],
                                  c["Linkedin"], c["Website"]) if x)
    if line:
        add("contact", line)
    add("note", "Career record: everything you've recorded, kept in full. "
                "A resume draws a selection from this.")

    def section(heading, items):
        if items:
            add("heading", heading, True)
            for kind, content in items:
                add(kind, content)

    if rec.summary:
        section("SUMMARY", [("text", rec.summary)])
    if rec.sets_apart:
        # Stored as the lines of one paragraph; blank lines separate paragraphs.
        paras = " ".join(x.strip() or "\n" for x in rec.sets_apart).split("\n")
        section("YOUR CAREER IN YOUR OWN WORDS", [("text", x.strip()) for x in paras if x.strip()])

    def jobs(outside):
        items = []
        for _, idx in journey._groups(rec):
            shown = [i for i in idx if rec.roles[i].outside() == outside]
            if not shown:
                continue
            items.append(("employer", rec.roles[shown[0]].employer or "[employer]"))
            for i in shown:
                role = rec.roles[i]
                items.append(("role", (role.title or "[title]",
                                       journey.clean(role.fields.get("Dates", "")))))
                for label, value in role.fields.items():
                    value = journey.clean(value)
                    if label != "Dates" and value:
                        items.append(("field", (label, value)))
                items.extend(("text", n) for n in role.notes if n.strip())
                for a in role.accomplishments:
                    if a.is_empty():
                        continue
                    items.append(("bullet", journey.line_text(a)))
                    for label, value in (("Situation", a.problem), ("What you did", a.actions),
                                         ("Your part", a.contribution), ("Result", a.results),
                                         ("Skills used", a.skills_used)):
                        if value and value.strip():
                            items.append(("detail", (label, value.strip())))
                if role.recorded_bullets:
                    items.append(("subhead", "From your resume"))
                    items.extend(("bullet", b) for b in role.recorded_bullets)
        return items

    for heading, items in (("EXPERIENCE", jobs(False)), ("WORK OUTSIDE PAID JOBS", jobs(True))):
        if not items:
            continue
        add("heading", heading, True)
        for n, (kind, content) in enumerate(items):
            nxt = items[n + 1][0] if n + 1 < len(items) else ""
            # Headings and labels hold on to what follows; a line holds on
            # to its own details so the story never leaves its line behind.
            keep = (kind in ("employer", "subhead")
                    or (kind in ("role", "field") and nxt in ("field", "text", "bullet", "subhead"))
                    or (kind == "bullet" and nxt == "detail")
                    or (kind == "detail" and nxt == "detail"))
            add(kind, content, keep)

    mine = [s for s in rec.skills if s.have == sk.YES]
    if mine:
        groups = {}
        for s in mine:
            groups.setdefault(s.category or "Other", []).append(
                s.name + (f" ({s.level})" if s.level else ""))
        section("SKILLS", [("field", (cat, ", ".join(names))) for cat, names in groups.items()])
    section("EDUCATION", [("text", x) for x in rec.education])
    section("LICENCES AND CERTIFICATIONS", [("text", x) for x in rec.certifications])
    for key, heading, _ in mr.EXTRA_SECTIONS:
        section(heading.upper(), [("text", x) for x in rec.extras.get(key, [])])
    return out


def master_docx(rec) -> bytes:
    """The career record as a Word document."""
    return _docx(master_paragraphs(rec))
